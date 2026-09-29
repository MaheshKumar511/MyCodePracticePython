from openpyxl import load_workbook


class ExcelUtils:

    @staticmethod
    def read_excel(file_path: str, sheet_name: str) -> list[dict]:
        workbook = load_workbook(file_path, data_only=True)
        sheet = workbook[sheet_name]

        rows = list(sheet.iter_rows(values_only=True))
        headers = rows[0]

        return [
            dict(zip(headers, row))
            for row in rows[1:]
            if any(value is not None for value in row)
        ]